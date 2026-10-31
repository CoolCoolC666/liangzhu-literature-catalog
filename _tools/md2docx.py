#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""md -> docx 转换器（为这份课程报告写的，不是通用 pandoc 替代品）。

为什么不用 pandoc：本机没装（`where pandoc` 空），而这份文档结构很固定，
用 python-docx 直接映射反而更可控，也少一个外部依赖。

处理到的 markdown 子集（够本报告用）：
    # / ## / ###        标题（自动挂中文字体，黑体）
    **加粗**             行内加粗
    `代码`               行内代码（等宽 + 深色）
    ```代码块```         围栏代码块（等宽小字，逐行成段）
    | a | b |           表格（首行当表头，Table Grid 边框）
    ---                 分隔线
    1. 2. 3.            有序列表（保留原文编号）
    - x                  无序列表
    *斜体*               行内斜体
    > 引用               引用块（左缩进 + 灰字）

中文排版要点（不设就会很难看）：
    python-docx 只设 font.name 时，Word 渲染中文会用默认的"等线"，
    跟正文混排时字号行距全乱。必须同时写 w:eastAsia。

用法：
    py -3 md2docx.py <输入.md> [-o 输出.docx]
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

BODY_CN = "宋体"
HEAD_CN = "黑体"
MONO = "Consolas"
BODY_PT = 11.0

# 行内标记：**粗** `码` *斜*
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+\*)")


def set_cn(run, ascii_font: str, cn_font: str, size: float | None = None,
           bold: bool | None = None, color: RGBColor | None = None) -> None:
    """python-docx 光设 font.name 不管中文，必须补 w:eastAsia。"""
    run.font.name = ascii_font
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn_font)


def add_inline(par, text: str, base_size: float = BODY_PT,
               bold: bool = False) -> None:
    """把一行 markdown 文本拆成若干带格式的 run 塞进段落。

    bold 必须往下传：`**`issue` 类关系`**` 这种加粗里嵌行内代码的写法，
    第一版只剥了外层 **，结果反引号原样留在 Word 正文里（被残留检查抓到）。
    """
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            add_inline(par, piece[2:-2], base_size, bold=True)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            set_cn(r, MONO, BODY_CN, base_size - 0.5, bold=bold,
                   color=RGBColor(0xA3, 0x1D, 0x1D))
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            set_cn(r, "Times New Roman", BODY_CN, base_size, bold=bold)
            r.font.italic = True
        else:
            r = par.add_run(piece)
            set_cn(r, "Times New Roman", BODY_CN, base_size, bold=bold)


def is_table_row(line: str) -> bool:
    return line.lstrip().startswith("|") and line.rstrip().endswith("|")


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_sep(line: str) -> bool:
    """markdown 表格的分隔行：|---|---| 或 |:--:|"""
    return is_table_row(line) and set(line.strip().replace("|", "").strip()) <= set("-: ")


def is_quote(line: str) -> bool:
    """markdown 引用块：> 开头"""
    return line.lstrip().startswith(">")


def quote_text(line: str) -> str:
    """剥掉引用块标记（> 与一个可选空格）"""
    return line.lstrip()[1:].lstrip()


def is_fence(line: str) -> bool:
    """markdown 围栏代码块：``` 开头"""
    return line.lstrip().startswith("```")


def add_inline_heading(par, text: str, size: float) -> None:
    """标题里也要解行内标记。

    ★ 以前标题是 `h.add_run(text)` 直接塞原文，于是
      `### 3.1 分词模式是精确模式的**子集**，不是超集` 在 Word 里
      会把 ** 子集** 原样显示出来（标题不参与 add_inline）。
    """
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            add_inline_heading(par, piece[2:-2], size)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            set_cn(r, MONO, HEAD_CN, size - 0.5, color=RGBColor(0xA3, 0x1D, 0x1D))
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            set_cn(r, "Times New Roman", HEAD_CN, size)
            r.font.italic = True
        else:
            r = par.add_run(piece)
            set_cn(r, "Times New Roman", HEAD_CN, size, bold=True)


def add_code_block(doc, lines: list[str]) -> int:
    """围栏代码块：等宽小字，逐行成段。

    ★ 以前没有这个分支，``` 围起来的整块会被当成**一个普通段落**，
      于是 JSON/SQL 源码在 Word 里挤成一坨，还带着 ``` 残留。
    """
    n = 0
    for raw in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(raw if raw else " ")
        set_cn(r, MONO, BODY_CN, BODY_PT - 1.5, color=RGBColor(0x22, 0x33, 0x44))
        n += 1
    # 块后留一个小空行
    doc.add_paragraph()
    return n


def convert(md_path: Path, out_path: Path) -> tuple[int, int, int]:
    lines = md_path.read_text(encoding="utf-8").splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.8)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(BODY_PT)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)

    n_para = n_table = 0
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 空行
        if not stripped:
            i += 1
            continue

        # 分隔线
        if stripped in ("---", "***", "___"):
            p = doc.add_paragraph()
            set_cn(p.add_run("─" * 34), "Times New Roman", BODY_CN, 9)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            h = doc.add_heading(level=level)
            sizes = {1: 18, 2: 15, 3: 12.5}
            add_inline_heading(h, text, sizes[level])
            h.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            h.paragraph_format.space_after = Pt(6)
            n_para += 1
            i += 1
            continue

        # 围栏代码块
        if is_fence(stripped):
            j = i + 1
            body: list[str] = []
            while j < len(lines) and not is_fence(lines[j]):
                body.append(lines[j].rstrip())
                j += 1
            j = min(j + 1, len(lines))          # 吃掉收尾的 ```
            n_para += add_code_block(doc, body)
            i = j
            continue

        # 表格（吃掉表头 + 分隔行 + 若干数据行）
        if is_table_row(stripped):
            header = split_row(stripped)
            j = i + 1
            if j < len(lines) and is_sep(lines[j].strip()):
                j += 1
            rows: list[list[str]] = []
            while j < len(lines) and is_table_row(lines[j].strip()):
                rows.append(split_row(lines[j].strip()))
                j += 1
            t = doc.add_table(rows=1 + len(rows), cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, txt in enumerate(header):
                cell = t.cell(0, c)
                cell.text = ""
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, txt, base_size=10)
                for r_ in p.runs:
                    r_.font.bold = True
            for ri, row in enumerate(rows, start=1):
                for c, txt in enumerate(row):
                    if c >= len(header):
                        continue
                    cell = t.cell(ri, c)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_inline(p, txt, base_size=10)
            doc.add_paragraph()
            n_table += 1
            i = j
            continue

        # 有序列表：保留原文编号（markdown 编号可能跳号，自己重排会失真）
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, f"{m.group(1)}. {m.group(2)}")
            n_para += 1
            i += 1
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, f"· {m.group(1)}")
            n_para += 1
            i += 1
            continue

        # 引用块：> 开头的行合并成一段（左缩进 + 灰字）
        # ★ 以前这里没有分支，"> xxx" 会被当普通段落原样写进 Word，
        #   于是正文里到处残留 "> " 符号（既有几份报告的 docx 都带着这个问题）。
        if is_quote(stripped):
            buf = [quote_text(stripped)]
            i += 1
            while i < len(lines) and lines[i].strip() and is_quote(lines[i].strip()):
                buf.append(quote_text(lines[i].strip()))
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, " ".join(x for x in buf if x), base_size=BODY_PT - 0.5)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x44, 0x55, 0x66)
            n_para += 1
            continue

        # 普通段落：把连续非空行并成一段（markdown 里软换行不该变成新段）
        buf = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r"^(#{1,3}\s|[-*]\s|\d+\.\s|\||>|```|---)", lines[i].strip()):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(buf))
        n_para += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return n_para, n_table, len(lines)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("-o", "--out", default=None)
    args = ap.parse_args()

    src = Path(args.src)
    out = Path(args.out) if args.out else src.with_suffix(".docx")
    if not src.exists():
        print(f"[FAIL] 找不到 {src}")
        return 1

    n_para, n_table, n_line = convert(src, out)
    print(f"[OK] {src.name} -> {out.name}")
    print(f"     源 {n_line} 行 -> 段落 {n_para} 个、表格 {n_table} 个")
    print(f"     大小 {out.stat().st_size:,} 字节")
    return 0


if __name__ == "__main__":
    sys.exit(main())
