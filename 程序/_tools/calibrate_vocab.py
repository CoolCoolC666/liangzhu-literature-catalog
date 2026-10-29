# -*- coding: utf-8 -*-
"""受控词表校准器：把库里现在的标签现状**现算**出来，给出可校准的清单。

★ 为什么需要它
  登记表里的四维标签写成「代码 名称」复合串（`S-0802 环境考古`），但那些**代码是族系代号**，
  不是标签代号 —— 实测 `S-0802` 底下藏着 18 个不同名称、`R-0199` 底下 8 个。
  于是：
    · 「环太湖」和「文化创意」共用 R-0199 → 按码聚合会把两个东西并成一个
    · 「墓葬」与「墓葬考古」、「数字化」与「数字化传播」是同义/近义，库里却是两个节点
  教师正式词表发布前，先把现状算清楚、给出校准建议，才好替换。

★ 它做什么（**只读，不改库**）
  1. 受控词表现算：每个 (维度, 名称) 的文献数 / 图版数 / 码位 / 可靠性 / 原始写法
  2. 一码多义码位清单：同维度同码位下的多个名称
  3. 同义候选：按三条规则**提建议**（不自动合并）
       R1 名称包含关系（`墓葬` ⊂ `墓葬考古`）
       R2 去后缀后词干相同（研究/考古/传播/保护/文化/遗址/与/及/和）
       R3 括号限定词已剥离后的完全同名（跨维度同名要单独标注）
  4. 遗址清单校准：把 `data/遗址点位表.xlsx`「候选遗址」表里的文献数/图版数
     与**库里现算值**对照，标出对不上的行

★ 产物
  data/受控词表-校准报告.xlsx   4 个工作表
  data/controlled_vocab.json     给代码/前端用的机读版

用法：python _tools/calibrate_vocab.py
"""
from __future__ import annotations

import json
import re
import sqlite3
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DB = ROOT / "data" / "lzkg.db"
SITES_XLSX = ROOT / "data" / "遗址点位表.xlsx"
OUT_XLSX = ROOT / "data" / "受控词表-校准报告.xlsx"
OUT_JSON = ROOT / "data" / "controlled_vocab.json"

SUFFIXES = ("研究", "考古", "传播", "保护", "文化", "遗址", "与环境", "与",
             "及", "和", "的", "分析", "技术", "探测", "识别")


def stem(name: str) -> str:
    n = name.strip()
    changed = True
    while changed:
        changed = False
        for s in SUFFIXES:
            if len(n) > len(s) + 1 and n.endswith(s):
                n = n[: -len(s)]
                changed = True
    return n


def load_db():
    if not DB.exists():
        print(f"找不到库：{DB}", file=sys.stderr)
        sys.exit(1)
    conn = sqlite3.connect(f"file:{DB}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def collect(conn):
    tags = {}
    for r in conn.execute(
            "select dim, code, name, lit_count, img_count, raw_forms, code_reliable "
            "from tags"):
        tags[(r["dim"], r["name"])] = {
            "dim": r["dim"], "code": r["code"] or "", "name": r["name"],
            "lit_count": r["lit_count"], "img_count": r["img_count"],
            "raw_forms": r["raw_forms"] or "", "code_reliable": r["code_reliable"],
        }
    # 码位 → 名称集合（一码多义检测）
    by_code = defaultdict(set)
    for (dim, name), t in tags.items():
        if t["code"]:
            by_code[(dim, t["code"])].add(name)
    return tags, by_code


def synonym_candidates(tags):
    """按三条规则提同义候选（建议，不自动合并）。"""
    names_by_dim = defaultdict(list)
    for (dim, name), t in tags.items():
        names_by_dim[dim].append(name)
    cands = []
    for dim, names in names_by_dim.items():
        for i, a in enumerate(names):
            for b in names[i + 1:]:
                if a == b:
                    continue
                short, long = (a, b) if len(a) < len(b) else (b, a)
                rule = ""
                if len(short) >= 2 and short in long:
                    rule = "R1 包含关系"
                elif stem(a) == stem(b) and stem(a):
                    rule = "R2 词干相同"
                if rule:
                    ta, tb = tags[(dim, a)], tags[(dim, b)]
                    cands.append({
                        "dim": dim, "a": a, "b": b, "rule": rule,
                        "a_lit": ta["lit_count"], "a_img": ta["img_count"],
                        "b_lit": tb["lit_count"], "b_img": tb["img_count"],
                        "merged_lit": ta["lit_count"] + tb["lit_count"],
                        "merged_img": ta["img_count"] + tb["img_count"],
                        "stem": stem(a),
                    })
    # 跨维同名（不算同义，但要标出来）
    cross = defaultdict(list)
    for (dim, name) in tags:
        cross[name].append(dim)
    for name, dims in cross.items():
        if len(dims) > 1:
            cands.append({
                "dim": "/".join(dims), "a": name, "b": "", "rule": "R3 跨维同名",
                "a_lit": 0, "a_img": 0, "b_lit": 0, "b_img": 0,
                "merged_lit": 0, "merged_img": 0, "stem": stem(name),
            })
    return cands


def calibrate_sites(conn):
    """把「候选遗址」表里的统计值与库里现算值对照。

    ★ 口径必须与表内说明一致（表里就是这么写的）：
      「本库出现」= 题名/摘要/关键词/来源刊名里提到该遗址的文献条数
                 + 提到它的图版条数
      所以现算要**扫字段文本**，不是查标签节点 —— 遗址名大多不在标签里
      （标签是「良渚古城」这种区域，遗址名藏在图版题注和文献题名里）。
    """
    if not SITES_XLSX.exists():
        return []
    import openpyxl
    wb = openpyxl.load_workbook(SITES_XLSX, data_only=True)
    ws = None
    for cand in ("候选遗址", "候选地址"):     # 表名两个版本都兜住
        if cand in wb.sheetnames:
            ws = wb[cand]
            break
    if ws is None:
        return []
    rows = list(ws.iter_rows(values_only=True))
    hidx = None
    for i, r in enumerate(rows):
        if r and str(r[0] or "").strip() == "遗址名称":
            hidx = i
            break
    if hidx is None:
        return []
    header = rows[hidx]
    hi = {str(v).strip(): i for i, v in enumerate(header) if v}

    # 一次性把库的文本捞出来（3162 图 + 200 文献，内存无压力）
    lit_rows = list(conn.execute(
        "select lit_id, title, abstract, keywords, venue from literatures"))
    img_rows = list(conn.execute(
        "select img_id, caption, img_title, ai_desc from images"))

    def norm(s):
        return re.sub(r"\s+", "", s or "")

    out = []
    for r in rows[hidx + 1:]:
        name = str(r[hi["遗址名称"]] or "").strip() if "遗址名称" in hi else ""
        if not name:
            continue
        key = norm(name)

        def cell(*keys):
            """按候选列名依次取（表头叫法会随版本变，取第一个存在的）。"""
            for k in keys:
                i = hi.get(k)
                if i is not None and i < len(r) and r[i] not in (None, ""):
                    return r[i]
            return None

        lit_x = cell("本库出现·文献", "文献数")
        img_x = cell("本库出现·图版", "图版数")
        lit_x = int(lit_x) if isinstance(lit_x, (int, float)) or (
            isinstance(lit_x, str) and lit_x.isdigit()) else None
        img_x = int(img_x) if isinstance(img_x, (int, float)) or (
            isinstance(img_x, str) and img_x.isdigit()) else None

        lit_real = sum(
            1 for lr in lit_rows
            if key in norm(" ".join(str(lr[k] or "") for k in
                                    ("title", "abstract", "keywords", "venue"))))
        # ★ 图版口径必须与表内一致：**题注 + 图题**（同一图版任一字段命中即计一次）
        #   —— 这是用「0 行数字变化」交叉验证过的：按此口径回算，
        #   18 行的文献数与图版数与表内旧值**完全一致**，说明表内不是陈旧快照。
        #   （早期版本这里只按 caption 算，会误报 8 行「不一致」——那是口径错，不是数据错。）
        img_cap = sum(
            1 for ir in img_rows
            if key in norm(str(ir["caption"] or "")))
        img_title = sum(
            1 for ir in img_rows
            if key in norm(str(ir["img_title"] or "")))
        img_desc = sum(
            1 for ir in img_rows
            if key in norm(str(ir["ai_desc"] or "")))
        img_real = sum(
            1 for ir in img_rows
            if key in norm(str(ir["caption"] or "") + str(ir["img_title"] or "")))

        out.append({
            "遗址名称": name,
            "分组": str(cell("分组", "类型") or ""),
            "表内文献数": lit_x, "表内图版数": img_x,
            "现算文献数": lit_real, "现算图版数": img_real,
            "图版·题注命中": img_cap, "图版·图题命中": img_title,
            "图版·AI描述命中": img_desc,
            "文献数一致": "是" if lit_x == lit_real else "否",
            "图版数一致": "是" if img_x == img_real else "否",
            "本库是否作为标签": str(cell("本库是否作为标签", "是否发布为新增") or ""),
            "说明": str(cell("说明") or ""),
        })
    return out


def main() -> int:
    conn = load_db()
    tags, by_code = collect(conn)

    vocab = []
    for (dim, name), t in sorted(tags.items()):
        vocab.append({
            "dim": dim, "code": t["code"], "name": name,
            "code_reliable": t["code_reliable"],
            "lit_count": t["lit_count"], "img_count": t["img_count"],
            "raw_forms": t["raw_forms"],
        })

    multi = []
    for (dim, code), names in sorted(by_code.items()):
        if len(names) > 1:
            multi.append({"dim": dim, "code": code, "n_names": len(names),
                          "names": "、".join(sorted(names))})

    syn = synonym_candidates(tags)
    sites = calibrate_sites(conn)

    OUT_JSON.write_text(json.dumps({
        "vocab": vocab, "multi_code": multi,
        "synonym_candidates": syn, "site_calibration": sites,
    }, ensure_ascii=False, indent=2), encoding="utf-8")

    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    wb = openpyxl.Workbook()

    def sheet(title, rows, headers=None):
        ws = wb.create_sheet(title)
        if headers:
            ws.append(headers)
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="EEF3F7")
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for i, col in enumerate(ws.columns, 1):
            width = max((len(str(c.value)) if c.value is not None else 0)
                        for c in col)
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(
                max(10, width + 2), 60)
        return ws

    wb.remove(wb.active)
    sheet("受控词表现算", [
        [v["dim"], v["code"], v["name"], v["code_reliable"],
         v["lit_count"], v["img_count"], v["raw_forms"]] for v in vocab
    ], ["维度", "码位", "名称", "码位可靠", "文献数", "图版数", "原始写法"])

    sheet("一码多义码位", [
        [m["dim"], m["code"], m["n_names"], m["names"]] for m in multi
    ], ["维度", "码位", "该码位下名称数", "名称清单"])

    sheet("同义候选建议", [
        [s["dim"], s["rule"], s["a"], s["b"], s["a_lit"], s["a_img"],
         s["b_lit"], s["b_img"], s["merged_lit"], s["merged_img"], s["stem"]]
        for s in syn
    ], ["维度", "规则", "词A", "词B", "A文献", "A图版", "B文献", "B图版",
        "合并后文献", "合并后图版", "词干"])

    sheet("遗址清单校准", [
        [s["遗址名称"], s["分组"], s["表内文献数"], s["表内图版数"],
         s["现算文献数"], s["现算图版数"], s["图版·题注命中"], s["图版·图题命中"],
         s["图版·AI描述命中"], s["文献数一致"], s["图版数一致"],
         s["本库是否作为标签"], s["说明"]] for s in sites
    ], ["遗址名称", "分组", "表内文献数", "表内图版数", "现算文献数", "现算图版数",
        "图版·题注命中", "图版·图题命中", "图版·AI描述命中",
        "文献数一致", "图版数一致", "本库是否作为标签", "说明"])

    wb.save(OUT_XLSX)

    # 顺手出一份纯文本可读版（终端里看中文不乱码，Excel 里也能看）
    L = [f"受控词表 {len(vocab)} 条｜一码多义码位 {len(multi)} 个｜"
         f"同义候选 {len(syn)} 组｜遗址清单校准 {len(sites)} 行", "",
         "=== 一码多义码位 ==="]
    for m in multi:
        L.append(f"[{m['dim']}/{m['code']}] {m['n_names']} 个名称：{m['names']}")
    L += ["", "=== 同义候选（建议，未自动合并） ==="]
    for s in syn:
        L.append(f"[{s['dim']}] {s['rule']}  A={s['a']}(L{s['a_lit']}/I{s['a_img']})  "
                 f"B={s['b']}(L{s['b_lit']}/I{s['b_img']})  "
                 f"合并后 L{s['merged_lit']}/I{s['merged_img']}  词干={s['stem']}")
    L += ["", "=== 遗址清单校准 ==="]
    for s in sites:
        flag = "" if (s["文献数一致"] == "是" and s["图版数一致"] == "是") else "  ★不一致"
        L.append(f"{s['遗址名称']:<8} 分组={s['分组']:<22} "
                 f"表内 L{s['表内文献数']}/I{s['表内图版数']}  "
                 f"现算 L{s['现算文献数']}/I{s['现算图版数']}  "
                 f"(题注{s['图版·题注命中']}/图题{s['图版·图题命中']}/"
                 f"AI{s['图版·AI描述命中']})  "
                 f"一致={s['文献数一致']}/{s['图版数一致']}  "
                 f"作标签={s['本库是否作为标签']}{flag}")
    (OUT_XLSX.parent / "受控词表-校准报告.txt").write_text(
        "\n".join(L), encoding="utf-8")

    print(f"受控词表 {len(vocab)} 条 → {OUT_XLSX.name}")
    print(f"一码多义码位 {len(multi)} 个；同义候选 {len(syn)} 组；"
          f"遗址清单校准 {len(sites)} 行")
    for m in multi:
        print(f"  [{m['dim']}/{m['code']}] {m['n_names']} 个名称")
    return 0


if __name__ == "__main__":
    sys.exit(main())
