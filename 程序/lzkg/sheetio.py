# -*- coding: utf-8 -*-
"""xlsx 表格读取：**按列名定位表头**，而不是按「第 4 行」。

★ 为什么必须拆成两套
--------------------
同一份「登记表」有两个用途，判据完全不同：

| 用途 | 表头在哪 | 判据 |
|---|---|---|
| **迁移**（读我们自己维护的最终版） | 固定在第 4 行 | 位置判据够用，但**按列名更稳** |
| **导入**（读别人填的表） | ★ **不确定** | ★ **必须按列名**：对方可能在表头前面多插两行说明，或把「填写说明」写成多列分列说明 |

用「非空单元格数 > 5」这种启发式定位表头，导入时整表会错位而且**不报错**。
所以这里统一成：**扫描前 12 行，找含关键列名的那一行**。

列名比对规则：
  1. 去掉尾部 `*`（必填标记不是列名的一部分）
  2. 去掉所有空白与全角/半角差异
  3. 两边都去 `（可选）` 之类的后缀说明
"""
from __future__ import annotations

import re
from pathlib import Path

from .registry import kind, norm_name, strip_star

# 「文献标识符*」这类带尾缀标记 —— 归一规则统一在 registry.norm_name 里
# （列定义侧与表头侧必须用**同一套**规则，否则「（可选）」类列名会对不上）


def norm(name: str) -> str:
    """列名归一：去尾 *、去（可选）、去所有空白与全角括号差异。

    ★ 直接复用 `registry.norm_name`：列定义的键也用它生成，
      两边规则不一致就会出现「表里明明有这列，却被当成未知列」。
    """
    return norm_name(name)


def find_header_row(rows: list[list], table: str) -> int:
    """在前 12 行里找表头行：含该表的主键列名（去 * 后）的那一行。

    找不到就抛错并说清楚「表里没有『文献标识符』这一列」——
    **宁可报错，也不要静默错位**。
    """
    spec = kind(table)
    pk = norm(spec["pk"])
    for i, r in enumerate(rows[:12]):
        names = [norm(v) for v in r if v not in (None, "")]
        if pk in names:
            return i
    raise KeyError(
        f"在前 12 行内找不到表头行：需要含「{spec['pk']}」这一列。"
        f"前 12 行各列名预览："
        + " | ".join(str(r[0])[:20] if r else "" for r in rows[:12]))


def read_rows(path: Path, table: str) -> tuple[list[dict], dict]:
    """读一张登记表 → (数据行 dict 列表, 表信息)。

    每行 dict 的 key 是**库字段名**（值取自列定义的 field），
    列名带 `*` 的不会被当成字段名的一部分。

    ★ meta 里额外给出**导入预检需要的判据**（不再另写一套表头定位逻辑）：
      - `header_raw`   表头原始字符串列表
      - `unknown_columns` 表里有、但列定义里没有的列（导入者自己加的列）
      - `missing_required` 列定义里必填、但表里没有的列
      - `variant`      FULL / MINIMAL —— 按表里有没有「完整版专属列」判定
        （★ 32 列模板没有「数据来源类型*」，它只对最终版必填；
          必填集合必须按变体裁剪，否则导入 32 列模板会被误判为「缺必填」）
    """
    import openpyxl
    from .registry import VARIANT_FULL, VARIANT_MINIMAL, columns, required_columns
    spec = kind(table)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    if not rows:
        return [], {}
    hidx = find_header_row(rows, table)
    header = rows[hidx]
    header_raw = [("" if v is None else str(v).strip()) for v in header]
    # 列名 → 库字段
    col2field = {}
    known_names = set()
    for i, v in enumerate(header):
        if v in (None, ""):
            continue
        c = spec["by_name"].get(norm(v))
        if c and c.get("field"):
            col2field[i] = c["field"]
            known_names.add(norm(v))
    unknown = [v for v in header_raw if v and norm(v) not in spec["by_name"]]
    out = []
    for r in rows[hidx + 1:]:
        if not any(v not in (None, "") for v in r):
            continue
        d = {}
        for i, field in col2field.items():
            if i < len(r):
                v = r[i]
                d[field] = "" if v is None else (str(v).strip() if not isinstance(v, (int, float)) else v)
        if any(str(v).strip() for v in d.values()):
            out.append(d)
    # 变体判定：完整版专属列在不在表里
    full_only = {norm(c["name"]) for c in columns(table, VARIANT_FULL)
                 if not c.get("in_template32", True)}
    variant = VARIANT_FULL if (full_only & known_names) else VARIANT_MINIMAL
    req = required_columns(table, variant)
    present = {norm(v) for v in header_raw if v}
    # 必填列用**登记表原始列名**（带 `*`）报出：`*` 是必填标记，报告里看得见
    missing_required = [c["name"] for c in columns(table, variant)
                         if c.get("required") and norm(c["name"]) not in present]
    meta = {"sheet": spec["sheet"], "title": spec["title"], "note": spec["note"],
            "header_row": hidx + 1, "n_cols": len([v for v in header if v not in (None, "")]),
            "n_rows": len(out),
            "columns": [norm(v) for v in header if v not in (None, "")],
            "header_raw": header_raw,
            "unknown_columns": unknown,
            "missing_required": missing_required,
            "variant": variant,
            "col2field": col2field}
    return out, meta


def read_rows_by_column(path: Path, table: str) -> tuple[list[dict], dict]:
    """同 `read_rows`，但每行的 key 是**列名**（去 `*`、归一化后）。

    ★ 导入预检要按**列名**做字段级 diff（「题名」旧值 vs 新值），
      而写库要按**库字段名**。所以这里两种视图都给出：
      行 dict 用列名做 key，`meta["col2field"]` 给出「列序 → 库字段」。
    """
    import openpyxl
    from .registry import VARIANT_FULL, VARIANT_MINIMAL, columns, required_columns
    spec = kind(table)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    if not rows:
        return [], {}
    hidx = find_header_row(rows, table)
    header = rows[hidx]
    header_raw = [("" if v is None else str(v).strip()) for v in header]
    # 列序 → (归一化列名, 库字段)
    col2field: dict[int, str] = {}
    col2name: dict[int, str] = {}
    known_names = set()
    for i, v in enumerate(header):
        if v in (None, ""):
            continue
        n = norm(v)
        c = spec["by_name"].get(n)
        if c:
            known_names.add(n)
            col2name[i] = n
            if c.get("field"):
                col2field[i] = c["field"]
    unknown = [v for v in header_raw if v and norm(v) not in spec["by_name"]]
    out = []
    for r in rows[hidx + 1:]:
        if not any(v not in (None, "") for v in r):
            continue
        d = {}
        for i, name in col2name.items():
            if i < len(r):
                v = r[i]
                d[name] = "" if v is None else (
                    str(v).strip() if not isinstance(v, (int, float)) else v)
        if any(str(v).strip() for v in d.values()):
            out.append(d)
    full_only = {norm(c["name"]) for c in columns(table, VARIANT_FULL)
                 if not c.get("in_template32", True)}
    variant = VARIANT_FULL if (full_only & known_names) else VARIANT_MINIMAL
    present = {norm(v) for v in header_raw if v}
    missing_required = [c["name"] for c in columns(table, variant)
                         if c.get("required") and norm(c["name"]) not in present]
    meta = {"sheet": spec["sheet"], "title": spec["title"], "note": spec["note"],
            "header_row": hidx + 1, "n_cols": len([v for v in header if v not in (None, "")]),
            "n_rows": len(out),
            "columns": [norm(v) for v in header if v not in (None, "")],
            "header_raw": header_raw,
            "unknown_columns": unknown,
            "missing_required": missing_required,
            "variant": variant,
            "col2field": col2field,
            "col2name": col2name}
    return out, meta
